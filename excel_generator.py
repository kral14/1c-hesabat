# -*- coding: utf-8 -*-
import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_1c_excel(report_data, config, output_path):
    """
    Generates an authentic 1C:Enterprise styled Excel (.xlsx) workbook
    matching the exact visual layout, colors, grouping levels (urven),
    outline collapsing brackets, numbers, and metadata of 'Товары на складах'.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары на складах"

    # Enable Excel Grouping / Outline Display (Urven)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.outlinePr.applyStyles = False
    try:
        ws.sheet_view.showOutlineSymbols = True
        ws.sheet_view.showGridLines = True
    except Exception:
        pass
    if hasattr(ws, "views") and ws.views and ws.views.sheetView:
        ws.views.sheetView[0].showOutlineSymbols = True
        ws.views.sheetView[0].showGridLines = True

    # 1C Palette Styles
    FONT_FAMILY = "Arial"
    
    font_title = Font(name=FONT_FAMILY, size=13, bold=True, color="002060")
    font_meta_bold = Font(name=FONT_FAMILY, size=8.5, bold=True, color="333333")
    font_meta_text = Font(name=FONT_FAMILY, size=8.5, bold=False, color="444444")
    
    font_th = Font(name=FONT_FAMILY, size=8.5, bold=True, color="000000")
    font_grp0 = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
    font_grp1 = Font(name=FONT_FAMILY, size=8.5, bold=False, color="000000")
    font_grp2 = Font(name=FONT_FAMILY, size=8, italic=True, color="444444")
    font_total = Font(name=FONT_FAMILY, size=9, bold=True, color="000000")
    font_red = Font(name=FONT_FAMILY, size=8.5, color="CC0000")
    font_red_bold = Font(name=FONT_FAMILY, size=9, bold=True, color="CC0000")

    # Fills
    fill_th = PatternFill(start_color="F5F2E3", end_color="F5F2E3", fill_type="solid") # 1C header yellow/cream
    fill_th_price = PatternFill(start_color="EEF2F8", end_color="EEF2F8", fill_type="solid") # soft blue for price
    fill_th_sum = PatternFill(start_color="FDF6E7", end_color="FDF6E7", fill_type="solid") # warm gold for sum
    fill_level0 = PatternFill(start_color="F2EFE2", end_color="F2EFE2", fill_type="solid") # Warehouse row beige
    fill_total = PatternFill(start_color="E4DFCC", end_color="E4DFCC", fill_type="solid") # Total row
    
    # Borders
    thin_border_side = Side(border_style="thin", color="A09C8D")
    double_border_side = Side(border_style="double", color="605C4D")
    
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_border_side)
    
    # Alignments
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    indicators = config.get("indicators", {})
    show_price = indicators.get("showPrice", True)
    show_sum = indicators.get("showSum", True)
    
    show_bc_unit = bool(indicators.get("barcodeUnit") or indicators.get("barcode_unit"))
    show_bc_box = bool(indicators.get("barcodeBox") or indicators.get("barcode_box"))
    show_bc_block = bool(indicators.get("barcodeBlock") or indicators.get("barcode_block"))
    if not (show_bc_unit or show_bc_box or show_bc_block):
        if indicators.get("barcode", True):
            show_bc_unit = True
            show_bc_box = True

    # Multi-Price Types
    price_types = config.get("priceTypes") or config.get("price_types")
    if not price_types or not isinstance(price_types, list):
        pt_single = config.get("priceType") or config.get("price_type", "20")
        price_types = [pt_single] if pt_single else ["20"]
    clean_pts = []
    for pt in price_types:
        if isinstance(pt, dict):
            pts = str(pt.get("name") or pt.get("code") or "").strip()
        else:
            pts = str(pt).strip()
        if pts and pts not in clean_pts and pts != "[object Object]":
            clean_pts.append(pts)
    price_types = clean_pts or ["20"]

    params = config.get("parameters", {})
    neg_red = params.get("negativeRed", True)

    # Active Quantity Columns
    qty_cols = []
    if indicators.get("qtyStart", True): qty_cols.append(("start_bal", "Начальный остаток"))
    if indicators.get("qtyIn", True): qty_cols.append(("in_qty", "Приход"))
    if indicators.get("qtyOut", True): qty_cols.append(("out_qty", "Расход"))
    if indicators.get("qtyEnd", True): qty_cols.append(("end_bal", "Конечный остаток"))
    if indicators.get("qtyTurnover", False): qty_cols.append(("turnover", "Оборот"))

    # Determine Column Map
    curr_c = 1
    col_code = curr_c; curr_c += 1
    col_art = curr_c; curr_c += 1

    col_bc_unit = curr_c if show_bc_unit else None
    if show_bc_unit: curr_c += 1

    col_bc_box = curr_c if show_bc_box else None
    if show_bc_box: curr_c += 1

    col_bc_block = curr_c if show_bc_block else None
    if show_bc_block: curr_c += 1

    col_title = curr_c; curr_c += 1
    col_entry = curr_c; curr_c += 1
    col_num = curr_c; curr_c += 1

    price_col_map = {}
    if show_price:
        for pt in price_types:
            price_col_map[pt] = curr_c
            curr_c += 1

    col_qty_start_idx = curr_c
    qty_col_map = {}
    for q_id, q_label in qty_cols:
        qty_col_map[q_id] = curr_c
        curr_c += 1
    col_qty_end_idx = curr_c - 1

    sum_col_map = {}
    if show_sum:
        for pt in price_types:
            sum_col_map[pt] = curr_c
            curr_c += 1

    total_cols = curr_c - 1

    # ----------------------------------------------------
    # 1. METADATA HEADER (Lines 1 to 6)
    # ----------------------------------------------------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    cell_t = ws.cell(row=1, column=1, value=report_data.get("report_title", "Товары на складах"))
    cell_t.font = font_title
    cell_t.alignment = align_left
    ws.row_dimensions[1].height = 24

    # Period
    p_start = config.get("startDate") or config.get("start_date", "01.09.2026")
    p_end = config.get("endDate") or config.get("end_date", "30.09.2026")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    ws.cell(row=2, column=1, value=f"Период: {p_start} - {p_end}").font = font_meta_bold

    # Indicators text
    ind_parts = []
    if show_price: ind_parts.append(f"Цена [{', '.join(price_types)}]")
    for _, q_label in qty_cols: ind_parts.append(q_label)
    if show_sum: ind_parts.append(f"Сумма [{', '.join(price_types)}]")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=total_cols)
    ws.cell(row=3, column=1, value=f"Показатели: Показатели({', '.join(ind_parts)});").font = font_meta_text

    # Groupings text
    row_grps = config.get("rowGroupings") or config.get("row_groupings", [])
    grp_list = [f"{g.get('field')} ({g.get('type', 'Элементы')})" for g in row_grps]
    if not grp_list: grp_list = ["Склад (Элементы)", "Номенклатура (Элементы)"]
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=total_cols)
    ws.cell(row=4, column=1, value=f"Группировки строк: {'; '.join(grp_list)};").font = font_meta_text

    # Filters text
    filters_list = config.get("filters", [])
    active_filters = [f for f in filters_list if f.get("active")]
    meta_row = 5
    if active_filters:
        f_parts = []
        for f in active_filters:
            op_label = "Равно"
            comp = f.get("comparison", "")
            if comp == "in_list": op_label = "В группе из списка"
            elif comp == "not_equal": op_label = "Не равно"
            elif comp == "contains": op_label = "Содержит"
            val = f.get("value", "")
            if isinstance(val, list):
                val = ", ".join(str(v) for v in val)
            f_parts.append(f"{f.get('field')} {op_label} {val}")
        ws.merge_cells(start_row=meta_row, start_column=1, end_row=meta_row, end_column=total_cols)
        ws.cell(row=meta_row, column=1, value=f"Отборы: {'; '.join(f_parts)}").font = font_meta_text
        meta_row += 1

    # Empty spacer row
    row_idx = meta_row + 1

    # ----------------------------------------------------
    # 2. TABLE HEADERS (2 Rows)
    # ----------------------------------------------------
    th_row1 = row_idx
    th_row2 = row_idx + 1

    grp_fields = [g.get("field") for g in row_grps if g.get("field")]
    col3_title = "\n".join(grp_fields) if grp_fields else "Склад\nНоменклатура"

    fixed_headers = [
        (col_code, "Код"),
        (col_art, "Артикул"),
    ]
    if col_bc_unit: fixed_headers.append((col_bc_unit, "Штрихкод\n(ədəd)"))
    if col_bc_box: fixed_headers.append((col_bc_box, "Штрихкод\n(qutu)"))
    if col_bc_block: fixed_headers.append((col_bc_block, "Штрихкод\n(blok)"))
    fixed_headers.extend([
        (col_title, col3_title),
        (col_entry, "С вход"),
        (col_num, "Номер")
    ])

    for col_i, title in fixed_headers:
        ws.merge_cells(start_row=th_row1, start_column=col_i, end_row=th_row2, end_column=col_i)
        c1 = ws.cell(row=th_row1, column=col_i, value=title)
        c1.font = font_th
        c1.fill = fill_th
        c1.alignment = align_center
        c1.border = border_cell
        ws.cell(row=th_row2, column=col_i).border = border_cell

    if show_price:
        for pt, col_p in price_col_map.items():
            ws.merge_cells(start_row=th_row1, start_column=col_p, end_row=th_row2, end_column=col_p)
            cp = ws.cell(row=th_row1, column=col_p, value=f"Цена\n({pt})")
            cp.font = font_th
            cp.fill = fill_th_price
            cp.alignment = align_center
            cp.border = border_cell
            ws.cell(row=th_row2, column=col_p).border = border_cell

    # Header Group: Количество
    if qty_cols:
        ws.merge_cells(start_row=th_row1, start_column=col_qty_start_idx, end_row=th_row1, end_column=col_qty_end_idx)
        cq = ws.cell(row=th_row1, column=col_qty_start_idx, value="Количество")
        cq.font = font_th
        cq.fill = fill_th
        cq.alignment = align_center
        for cc in range(col_qty_start_idx, col_qty_end_idx + 1):
            ws.cell(row=th_row1, column=cc).border = border_cell

        # Row 2 Subheaders
        for q_id, q_label in qty_cols:
            c_idx = qty_col_map[q_id]
            c2 = ws.cell(row=th_row2, column=c_idx, value=q_label)
            c2.font = font_th
            c2.fill = fill_th
            c2.alignment = align_center
            c2.border = border_cell

    if show_sum:
        for pt, col_s in sum_col_map.items():
            ws.merge_cells(start_row=th_row1, start_column=col_s, end_row=th_row2, end_column=col_s)
            cs = ws.cell(row=th_row1, column=col_s, value=f"Сумма\n({pt})")
            cs.font = font_th
            cs.fill = fill_th_sum
            cs.alignment = align_center
            cs.border = border_cell
            ws.cell(row=th_row2, column=col_s).border = border_cell

    ws.row_dimensions[th_row1].height = 24
    ws.row_dimensions[th_row2].height = 20

    row_idx = th_row2 + 1

    # Styles for doc movement rows in Excel
    fill_doc = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")  # warm light orange
    font_doc = Font(name=FONT_FAMILY, size=8, italic=True, color="6B4C00")
    border_doc_top = Side(border_style="dashed", color="D4B88C")
    border_doc = Border(
        left=Side(border_style="thin", color="C8A97B"),
        right=thin_border_side,
        top=border_doc_top,
        bottom=border_doc_top
    )

    # ----------------------------------------------------
    # 3. DATA ROWS & OUTLINE GROUPING LEVELS (Urven)
    # ----------------------------------------------------
    items = report_data.get("items", [])
    num_fmt_qty = "#,##0.00"
    num_fmt_price = "#,##0.00"

    for item in items:
        level = item.get("level", 0)
        is_doc = item.get("is_doc", False)
        ws.row_dimensions[row_idx].height = 17 if is_doc else 19

        if level > 0:
            ws.row_dimensions[row_idx].outline_level = level

        # Font & Fill per level
        if is_doc:
            f_row = font_doc
            fill_row = fill_doc
            b_row = border_doc
            indent_spaces = "         "
        elif level == 0:
            f_row = font_grp0
            fill_row = fill_level0
            b_row = border_cell
            indent_spaces = ""
        elif level == 1:
            f_row = font_grp1
            fill_row = None
            b_row = border_cell
            indent_spaces = "   "
        else:
            f_row = font_grp2
            fill_row = None
            b_row = border_cell
            indent_spaces = "      "

        # Code & Artikul
        c_code = ws.cell(row=row_idx, column=col_code, value=item.get("code", "") or "")
        c_code.font = f_row; c_code.alignment = align_center; c_code.border = b_row
        if fill_row: c_code.fill = fill_row

        c_art = ws.cell(row=row_idx, column=col_art, value=item.get("artikul", "") or "")
        c_art.font = f_row; c_art.alignment = align_center; c_art.border = b_row
        if fill_row: c_art.fill = fill_row

        # Barcode cells
        if col_bc_unit:
            c_bcu = ws.cell(row=row_idx, column=col_bc_unit, value=item.get("barcode_unit", "") or "")
            c_bcu.font = f_row; c_bcu.alignment = align_center; c_bcu.border = b_row
            if fill_row: c_bcu.fill = fill_row
        if col_bc_box:
            c_bcb = ws.cell(row=row_idx, column=col_bc_box, value=item.get("barcode_box", "") or "")
            c_bcb.font = f_row; c_bcb.alignment = align_center; c_bcb.border = b_row
            if fill_row: c_bcb.fill = fill_row
        if col_bc_block:
            c_bck = ws.cell(row=row_idx, column=col_bc_block, value=item.get("barcode_block", "") or "")
            c_bck.font = f_row; c_bck.alignment = align_center; c_bck.border = b_row
            if fill_row: c_bck.fill = fill_row

        # Title with indent; for doc rows append date if available
        raw_title = item.get("title") or item.get("name", "")
        doc_date = item.get("doc_date", "") or ""
        if is_doc and doc_date:
            title_val = indent_spaces + str(raw_title) + "  [" + doc_date + "]"
        else:
            title_val = indent_spaces + str(raw_title)
        c_tit = ws.cell(row=row_idx, column=col_title, value=title_val)
        c_tit.font = f_row; c_tit.alignment = align_left; c_tit.border = b_row
        if fill_row: c_tit.fill = fill_row

        # Doc entry & number
        c_ent = ws.cell(row=row_idx, column=col_entry, value=item.get("custom_field", "") or "")
        c_ent.font = f_row; c_ent.alignment = align_center; c_ent.border = b_row
        if fill_row: c_ent.fill = fill_row

        doc_num_val = item.get("doc_number", "") or ""
        c_docnum = ws.cell(row=row_idx, column=col_num, value=doc_num_val)
        c_docnum.font = f_row; c_docnum.alignment = align_center; c_docnum.border = b_row
        if fill_row: c_docnum.fill = fill_row

        # Multi-Prices
        if show_price:
            for pt, col_p in price_col_map.items():
                p_val = item.get("prices", {}).get(pt) if isinstance(item.get("prices"), dict) else item.get("unit_price", 0.0)
                c_p = ws.cell(row=row_idx, column=col_p)
                if p_val and float(p_val) > 0:
                    c_p.value = float(p_val)
                    c_p.number_format = num_fmt_price
                else:
                    c_p.value = ""
                c_p.font = f_row; c_p.alignment = align_right; c_p.border = b_row
                if fill_row: c_p.fill = fill_row

        # Quantities
        for q_id, _ in qty_cols:
            c_idx = qty_col_map[q_id]
            val = item.get(q_id, 0.0)
            c_q = ws.cell(row=row_idx, column=c_idx)
            if val is not None:
                c_q.value = float(val)
                c_q.number_format = num_fmt_qty
                if neg_red and float(val) < 0:
                    c_q.font = font_red_bold if level == 0 else font_red
                else:
                    c_q.font = f_row
            c_q.alignment = align_right
            c_q.border = b_row
            if fill_row: c_q.fill = fill_row

        # Multi-Sums
        if show_sum:
            for pt, col_s in sum_col_map.items():
                s_val = item.get("sums", {}).get(pt) if isinstance(item.get("sums"), dict) else item.get("end_sum", 0.0)
                c_s = ws.cell(row=row_idx, column=col_s)
                if s_val is not None:
                    c_s.value = float(s_val)
                    c_s.number_format = num_fmt_price
                    if neg_red and float(s_val) < 0:
                        c_s.font = font_red_bold if level == 0 else font_red
                    else:
                        c_s.font = f_row
                c_s.alignment = align_right; c_s.border = b_row
                if fill_row: c_s.fill = fill_row

        row_idx += 1

    # ----------------------------------------------------
    # 4. GRAND TOTAL ROW (Итого)
    # ----------------------------------------------------
    totals = report_data.get("totals", {})
    ws.row_dimensions[row_idx].height = 22
    
    for c_i in range(1, total_cols + 1):
        cell_tot = ws.cell(row=row_idx, column=c_i)
        cell_tot.font = font_total
        cell_tot.fill = fill_total
        cell_tot.border = border_total

    ws.cell(row=row_idx, column=col_title, value=f"Итого ({len(items)} sətir)").alignment = align_left

    for q_id, _ in qty_cols:
        c_idx = qty_col_map[q_id]
        val = totals.get(q_id, 0.0)
        c_tot_q = ws.cell(row=row_idx, column=c_idx, value=float(val))
        c_tot_q.number_format = num_fmt_qty
        c_tot_q.alignment = align_right

    if show_sum:
        tot_sums_dict = totals.get("total_sums", {})
        for pt, col_s in sum_col_map.items():
            val_sum = tot_sums_dict.get(pt, totals.get("total_sum", 0.0))
            c_tot_s = ws.cell(row=row_idx, column=col_s, value=float(val_sum or 0.0))
            c_tot_s.number_format = num_fmt_price
            c_tot_s.alignment = align_right

    # ----------------------------------------------------
    # 5. COLUMN WIDTHS & FREEZE PANES
    # ----------------------------------------------------
    ws.column_dimensions[get_column_letter(col_code)].width = 12
    ws.column_dimensions[get_column_letter(col_art)].width = 15
    if col_bc_unit: ws.column_dimensions[get_column_letter(col_bc_unit)].width = 16
    if col_bc_box: ws.column_dimensions[get_column_letter(col_bc_box)].width = 16
    if col_bc_block: ws.column_dimensions[get_column_letter(col_bc_block)].width = 16
    ws.column_dimensions[get_column_letter(col_title)].width = 46
    ws.column_dimensions[get_column_letter(col_entry)].width = 9
    ws.column_dimensions[get_column_letter(col_num)].width = 13
    
    for pt, col_p in price_col_map.items():
        ws.column_dimensions[get_column_letter(col_p)].width = 14
    for q_id, _ in qty_cols:
        ws.column_dimensions[get_column_letter(qty_col_map[q_id])].width = 16
    for pt, col_s in sum_col_map.items():
        ws.column_dimensions[get_column_letter(col_s)].width = 18

    # Freeze header rows
    freeze_col = get_column_letter(col_title)
    ws.freeze_panes = f"{freeze_col}{th_row2 + 1}"

    # Save to output file
    dir_name = os.path.dirname(output_path)
    if dir_name:
        os.makedirs(dir_name, exist_ok=True)
    wb.save(output_path)
    print(f"[EXCEL EXPORT HAZIRDIR] {output_path} fayli ugurla yazildi (Setir: {len(items)}).", flush=True)
    return True


def generate_portfolio_catalog_excel(items, filters, output_path):
    """
    Generates an authentic 1C:Enterprise styled Excel workbook for the
    Portfolio Catalog (Реестр номенклатуры по портфелям) with all fields:
    Code, Artikul, CV Code, Barcode, Name, Folder, Nom Group, Portfolio,
    Item Type, Base Unit, Price, Manufacturer.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары по портфелям"

    try:
        ws.sheet_view.showGridLines = True
    except Exception:
        pass

    FONT_FAMILY = "Arial"
    font_title = Font(name=FONT_FAMILY, size=13, bold=True, color="002060")
    font_meta = Font(name=FONT_FAMILY, size=8.5, bold=False, color="444444")
    font_th = Font(name=FONT_FAMILY, size=8.5, bold=True, color="000000")
    font_td = Font(name=FONT_FAMILY, size=8.5, color="000000")
    font_code = Font(name="Consolas", size=8.5, color="003366")
    font_price = Font(name=FONT_FAMILY, size=8.5, bold=True, color="000080")

    fill_th = PatternFill(start_color="F5F2E3", end_color="F5F2E3", fill_type="solid")
    fill_zebra = PatternFill(start_color="FAF9F5", end_color="FAF9F5", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="A09C8D")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_th = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1. Title
    ws["B2"] = "1С:Предприятие — Реестр номенклатуры по портфелям"
    ws["B2"].font = font_title

    # 2. Metadata / Filter info
    ports_val = filters.get("portfolios") or filters.get("portfolio")
    if isinstance(ports_val, list):
        port_text = ", ".join(ports_val) if ports_val else "Все портфели"
    else:
        port_text = ports_val or "Все портфели"

    grps_val = filters.get("nom_groups") or filters.get("nom_group")
    if isinstance(grps_val, list):
        grp_text = ", ".join(grps_val) if grps_val else "Все группы"
    else:
        grp_text = grps_val or "Все группы"

    raw_pts = filters.get("price_types")
    if not raw_pts or not isinstance(raw_pts, list):
        single_pt = filters.get("price_type")
        raw_pts = [single_pt] if single_pt else []
    sel_pts = [str(p).strip() for p in raw_pts if str(p).strip() and str(p).lower() != "none" and str(p) != "(Без цен)"]
    price_text = ", ".join(sel_pts) if sel_pts else "Без цен"
    now_str = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")

    extra_notes = []
    if filters.get("selected_only"):
        extra_notes.append("Выбранные позиции")
    filter_note = f"  |  Фильтр: {', '.join(extra_notes)}" if extra_notes else ""

    ws["B3"] = f"Портфель: {port_text}  |  Группа: {grp_text}  |  Тип цен: {price_text}{filter_note}  |  Дата формирования: {now_str}  |  Всего: {len(items)} товаров"
    ws["B3"].font = font_meta

    # 3. Table Headers
    headers = [
        ("№", 5, align_center),
        ("Код", 12, align_center),
        ("Артикул", 14, align_left),
        ("СВ код", 12, align_center),
        ("Штрихкод", 16, align_center),
        ("Наименование", 38, align_left),
        ("Папка (Родитель)", 20, align_left),
        ("Номенклатурная группа", 20, align_left),
        ("Портфель", 18, align_left),
        ("Вид номенклатуры", 14, align_left),
        ("Базовая ед.", 10, align_center),
    ]
    if sel_pts:
        for pt in sel_pts:
            headers.append((f"Цена ({pt})", 13, align_right))
    else:
        headers.append(("Цена", 12, align_right))
    headers.append(("Производитель", 18, align_left))
    headers.append(("Комментарий", 25, align_left))

    header_row = 5
    ws.row_dimensions[header_row].height = 24

    for col_idx, (th_text, width, _) in enumerate(headers, start=2): # Start from col B
        cell = ws.cell(row=header_row, column=col_idx, value=th_text)
        cell.font = font_th
        cell.fill = fill_th
        cell.border = border_cell
        cell.alignment = align_th
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    analysis_info = filters.get("price_analysis") or {}
    analysis_enabled = analysis_info.get("enabled", False)
    base_pt = analysis_info.get("basePriceType", "20")

    fill_diff = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    font_diff = Font(name=FONT_FAMILY, size=8.5, bold=True, color="856404")
    fill_base = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    font_base = Font(name=FONT_FAMILY, size=8.5, bold=True, color="002060")

    # 4. Data Rows
    current_row = header_row + 1
    for i, itm in enumerate(items, start=1):
        ws.row_dimensions[current_row].height = 18
        is_even = (i % 2 == 0)

        row_vals = [
            (i, align_center, font_td, None, None),
            (itm.get("code") or "", align_center, font_code, None, None),
            (itm.get("artikul") or "", align_left, font_td, None, None),
            (itm.get("cv") or "", align_center, font_code, None, None),
            (itm.get("barcode") or "", align_center, font_td, "@", None),
            (itm.get("name") or "", align_left, font_td, None, None),
            (itm.get("folder") or "", align_left, font_td, None, None),
            (itm.get("group") or "", align_left, font_td, None, None),
            (itm.get("portfolio") or "", align_left, font_td, None, None),
            (itm.get("type") or "", align_left, font_td, None, None),
            (itm.get("unit") or "", align_center, font_td, None, None),
        ]
        prices_dict = itm.get("prices") or {}
        base_val = float(prices_dict.get(base_pt, itm.get("price", 0)) or 0) if analysis_enabled else 0.0

        if sel_pts:
            for pt in sel_pts:
                p_val = float(prices_dict.get(pt, itm.get("price", 0)) or 0)
                is_diff = analysis_enabled and (pt != base_pt) and (abs(p_val - base_val) > 0.001)
                is_benchmark = analysis_enabled and (pt == base_pt)

                f_price = font_diff if is_diff else (font_base if is_benchmark else font_price)
                fill_price = fill_diff if is_diff else (fill_base if is_benchmark else None)

                row_vals.append((p_val, align_right, f_price, "#,##0.00", fill_price))
        else:
            row_vals.append((float(itm.get("price") or 0), align_right, font_price, "#,##0.00", None))
        row_vals.append((itm.get("manufacturer") or "", align_left, font_td, None, None))
        row_vals.append((itm.get("comment") or "", align_left, font_td, None, None))

        for col_idx, item_tuple in enumerate(row_vals, start=2):
            val = item_tuple[0]
            alignment = item_tuple[1]
            font = item_tuple[2]
            num_fmt = item_tuple[3]
            c_fill = item_tuple[4] if len(item_tuple) > 4 else None

            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font
            cell.alignment = alignment
            cell.border = border_cell
            if num_fmt:
                cell.number_format = num_fmt
            if c_fill:
                cell.fill = c_fill
            elif is_even:
                cell.fill = fill_zebra

        current_row += 1

    # Freeze panes below header
    ws.freeze_panes = f"C{header_row + 1}"

    dir_name = os.path.dirname(output_path)
    if dir_name:
        os.makedirs(dir_name, exist_ok=True)
    wb.save(output_path)
    print(f"[PORTFOLIO EXCEL SAVED] {output_path} with {len(items)} items.", flush=True)
    return True



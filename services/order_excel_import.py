"""Read store quantities from Excel; never calls 1C or writes documents."""
import io
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook


def read_store_quantities(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        code_column, quantity_column = 0, 1
        records = []
        first = True
        code_headers = {'mağaza kodu', 'magaza kodu', 'kod', 'код магазина', 'код', 'store code'}
        quantity_headers = {'miqdar', 'say', 'количество', 'quantity', 'qty'}
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            values = [cell.value for cell in cells]
            if all(value is None or str(value).strip() == '' for value in values):
                continue
            if first:
                first = False
                labels = [str(value or '').strip().lower() for value in values]
                code_matches = [i for i, label in enumerate(labels) if label in code_headers]
                quantity_matches = [i for i, label in enumerate(labels) if label in quantity_headers]
                if code_matches and quantity_matches:
                    code_column, quantity_column = code_matches[0], quantity_matches[0]
                    continue
            if len(cells) <= max(code_column, quantity_column):
                raise ValueError(f'{row_number}-ci sətirdə mağaza kodu və miqdar olmalıdır.')
            raw_code = cells[code_column].value
            code = str(raw_code or '').strip()
            if isinstance(raw_code, (int, float)) and not isinstance(raw_code, bool):
                code = str(int(raw_code)) if float(raw_code).is_integer() else str(raw_code)
                fmt = cells[code_column].number_format
                if fmt and set(fmt) == {'0'}:
                    code = code.zfill(len(fmt))
            if not code or any(char.isspace() for char in code):
                raise ValueError(f'{row_number}-ci sətirdə mağaza kodu düzgün deyil.')
            try:
                qty = Decimal(str(cells[quantity_column].value).strip().replace(',', '.'))
                if not qty.is_finite() or qty <= 0 or qty != qty.quantize(Decimal('0.001')):
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                raise ValueError(f'{row_number}-ci sətirdə miqdar müsbət və ən çox 3 onluq rəqəmli olmalıdır.')
            records.append({'code': code, 'quantity': format(qty, 'f')})
        if not records:
            raise ValueError('Excel faylında mağaza kodu və miqdar tapılmadı.')
        return {'sheet': sheet.title, 'rows': records}
    finally:
        workbook.close()
